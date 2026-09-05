"""Mesin perhitungan simulasi Insentif Service MFlash.

Dipisah dari tampilan supaya mudah diuji dan dipakai ulang oleh aplikasi
pengajuan insentif bila skema ini jadi dipakai.
"""
from __future__ import annotations

import re
import unicodedata
from collections import defaultdict

from openpyxl import load_workbook

# Kolom yang harus ada pada ekspor "Rincian Faktur Penjualan".
KOLOM_WAJIB = ["NO FAKTUR", "KATEGORI BARANG", "TOTAL HARGA", "KATEGORI PENJUALAN"]

BUKAN_NAMA = {"", "N/A", "NA", "-", "NULL", "NONE", "0"}


def kunci(teks) -> str:
    """Normalkan nama supaya beda huruf besar/tanda baca tetap dianggap sama."""
    t = unicodedata.normalize("NFKD", str(teks or ""))
    return re.sub(r"[^a-z0-9]", "", t.lower())


def ambil(baris: dict, *nama_kolom):
    """Ambil nilai kolom; nama header berbeda antar versi ekspor."""
    for n in nama_kolom:
        v = baris.get(kunci(n))
        if v not in (None, ""):
            return v
    return None


def angka(v) -> float:
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).replace("Rp", "").replace(" ", "").replace(".", "").replace(",", ".")
    try:
        return float(t)
    except ValueError:
        return 0.0


def _isi_sheet(ws, batas=None):
    """Baca satu sheet.

    Sebagian ekspor menulis <dimension ref="A1"/> yang salah sehingga mode
    read_only hanya melihat satu baris; ukurannya dikosongkan lebih dulu.
    """
    ws._max_row = None
    ws._max_column = None
    keluar = []
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        keluar.append(r)
        if batas is not None and i + 1 >= batas:
            break
    return keluar


def baca_faktur(sumber, wajib=None, batas_header: int = 12) -> list[dict]:
    """Baca ekspor Excel menjadi list-of-dict, menelusuri semua sheet."""
    wajib = wajib or KOLOM_WAJIB
    wb = load_workbook(sumber, data_only=True, read_only=True)
    try:
        for nama in wb.sheetnames:
            awal = _isi_sheet(wb[nama], batas_header)
            idx = None
            for i, r in enumerate(awal):
                ada = {kunci(c) for c in r if c is not None}
                if all(kunci(w) in ada for w in wajib):
                    idx = i
                    break
            if idx is None:
                continue
            rows = _isi_sheet(wb[nama])
            header = [kunci(c) for c in rows[idx]]
            hasil = []
            for r in rows[idx + 1:]:
                if all(c is None or str(c).strip() == "" for c in r):
                    continue
                hasil.append({h: v for h, v in zip(header, r) if h})
            return hasil
    finally:
        wb.close()
    raise ValueError(
        "Kolom " + ", ".join(wajib) + " tidak ditemukan di berkas. "
        "Pastikan yang diunggah adalah ekspor Rincian Faktur Penjualan.")


def bulan_dari(v):
    try:
        return v.month
    except AttributeError:
        try:
            return int(str(v).split("-")[1])
        except (IndexError, ValueError):
            return None


def tahun_dari(v):
    return getattr(v, "year", None)


def periode_tersedia(faktur: list[dict]) -> list[tuple[int, int]]:
    """Daftar (tahun, bulan) yang ada di berkas, terbaru lebih dulu."""
    ada = set()
    for f in faktur:
        tgl = ambil(f, "TGL FAKTUR", "Tanggal Faktur", "Tanggal")
        b, t = bulan_dari(tgl), tahun_dari(tgl)
        if b:
            ada.add((t or 0, b))
    return sorted(ada, reverse=True)


def hitung(faktur: list[dict], bulan: int, tahun: int | None = None, *,
           pct_teknisi: float = 30.0,
           hanya_member: bool = False,
           pct_pool: float = 2.0,
           porsi: dict[str, float] | None = None,
           nama_store_leader: str = "Store Leader",
           pct_lama_team: float = 2.0) -> dict:
    """Hitung pool insentif service dan pembagiannya ke tiga peran.

    Alur:
      1. Omset jasa service dikumpulkan per faktur (kategori barang JASA).
      2. Bagi hasil MFlash = omset jasa x (100% - bagi hasil teknisi).
      3. Pool insentif  = bagi hasil MFlash x pct_pool.
      4. Pool dibagi ke Sales / Admin / Store Leader menurut porsi.
      5. Bagian Sales dan Admin dipecah pro-rata menurut kontribusi
         omset jasa masing-masing orang; bagian Store Leader utuh.
    """
    porsi = porsi or {"sales": 50.0, "admin": 30.0, "store_leader": 20.0}
    total_porsi = sum(porsi.values()) or 1.0

    omset_jasa = 0.0
    omset_jasa_member = 0.0
    omset_sparepart_service = 0.0
    per_sales: dict[str, float] = defaultdict(float)
    per_admin: dict[str, float] = defaultdict(float)
    asli_sales: dict[str, str] = {}
    asli_admin: dict[str, str] = {}
    tanpa_sales = 0.0
    tanpa_admin = 0.0
    faktur_service: set = set()
    dipakai = 0

    for f in faktur:
        tgl = ambil(f, "TGL FAKTUR", "Tanggal Faktur", "Tanggal")
        if bulan_dari(tgl) != bulan:
            continue
        th = tahun_dari(tgl)
        if tahun and th and th != tahun:
            continue
        dipakai += 1

        kat_barang = str(ambil(f, "Kategori Barang") or "").strip().upper()
        kat_jual = str(ambil(f, "Kategori Penjualan") or "").strip().upper()
        nilai = angka(ambil(f, "Total Harga"))
        member = kunci(ambil(f, "Kategori Pelanggan")) == kunci("MEMBER REGULER")

        if kat_jual.startswith("SERVICE"):
            faktur_service.add(str(ambil(f, "No Faktur") or ""))
            if kat_barang == "SPAREPART":
                omset_sparepart_service += nilai

        if kat_barang != "JASA":
            continue
        if hanya_member and not member:
            continue

        omset_jasa += nilai
        if member:
            omset_jasa_member += nilai

        nama_s = str(ambil(f, "Yang Menyerahkan/Menjual Faktur Penjualan",
                           "Yang Menyerahkan/Menjual") or "").strip()
        if nama_s.upper() in BUKAN_NAMA:
            tanpa_sales += nilai
        else:
            per_sales[kunci(nama_s)] += nilai
            asli_sales.setdefault(kunci(nama_s), nama_s)

        nama_a = str(ambil(f, "Nama Admin", "Admin") or "").strip()
        if nama_a.upper() in BUKAN_NAMA:
            tanpa_admin += nilai
        else:
            per_admin[kunci(nama_a)] += nilai
            asli_admin.setdefault(kunci(nama_a), nama_a)

    bagi_hasil = omset_jasa * (100.0 - pct_teknisi) / 100.0
    pool = bagi_hasil * pct_pool / 100.0
    bagian = {k: pool * v / total_porsi for k, v in porsi.items()}

    def pecah(per_orang, asli, jatah):
        total = sum(per_orang.values())
        baris = []
        for k, omset in sorted(per_orang.items(), key=lambda x: -x[1]):
            andil = (omset / total * 100.0) if total else 0.0
            baris.append({"nama": asli[k], "omset_jasa": round(omset),
                          "andil_pct": round(andil, 2),
                          "insentif": round(jatah * omset / total) if total else 0})
        return baris, total

    baris_sales, total_sales = pecah(per_sales, asli_sales, bagian["sales"])
    baris_admin, total_admin = pecah(per_admin, asli_admin, bagian["admin"])

    # Skema lama sebagai pembanding: 2% dari bagi hasil jasa member reguler.
    bagi_hasil_member = omset_jasa_member * (100.0 - pct_teknisi) / 100.0
    lama = round(bagi_hasil_member * pct_lama_team / 100.0)

    return {
        "bulan": bulan, "tahun": tahun,
        "jumlah_baris_diproses": dipakai,
        "jumlah_faktur_service": len(faktur_service),
        "omset_jasa": round(omset_jasa),
        "omset_jasa_member": round(omset_jasa_member),
        "omset_sparepart_service": round(omset_sparepart_service),
        "pct_teknisi": pct_teknisi,
        "bagi_hasil_mflash": round(bagi_hasil),
        "pct_pool": pct_pool,
        "pool": round(pool),
        "porsi": porsi,
        "bagian": {k: round(v) for k, v in bagian.items()},
        "sales": baris_sales,
        "admin": baris_admin,
        "store_leader": {"nama": nama_store_leader,
                         "insentif": round(bagian["store_leader"])},
        "omset_tanpa_sales": round(tanpa_sales),
        "omset_tanpa_admin": round(tanpa_admin),
        "skema_lama": lama,
        "selisih": round(pool) - lama,
    }


def nama_mirip(*kumpulan) -> list[tuple[str, str]]:
    """Pasangan nama yang satu memuat yang lain — kemungkinan orang yang sama
    dengan dua ejaan (mis. 'ARIF FIKRI' dan 'ARIF FIKRI KLENDER')."""
    semua = sorted({n for k in kumpulan for n in k})
    pasangan = []
    for i, a in enumerate(semua):
        for b in semua[i + 1:]:
            ka, kb = kunci(a), kunci(b)
            if ka != kb and (ka in kb or kb in ka):
                pasangan.append((a, b))
    return pasangan


def pct_netral_biaya(hasil: dict) -> float:
    """Tarif pool yang membuat biaya sama dengan skema lama."""
    dasar = hasil["bagi_hasil_mflash"]
    return round(hasil["skema_lama"] / dasar * 100.0, 3) if dasar else 0.0
